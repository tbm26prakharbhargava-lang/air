from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


SOURCE_MD = Path("/workspace/sarvam_ai_brand_teardown.md")
OUTPUT_PPTX = Path("/workspace/sarvam_ai_brand_teardown.pptx")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(11, 15, 20)
BG_SOFT = RGBColor(18, 25, 33)
ACCENT = RGBColor(255, 122, 26)
ACCENT_2 = RGBColor(0, 201, 167)
TEXT = RGBColor(245, 247, 250)
MUTED = RGBColor(191, 201, 212)
LINE = RGBColor(42, 51, 61)


def split_slides(markdown_text: str) -> list[str]:
    slides = [section.strip() for section in markdown_text.split("\n---\n")]
    return [slide for slide in slides if slide]


def parse_slide(section: str) -> tuple[str, list[str]]:
    lines = [line.rstrip() for line in section.splitlines()]
    title = ""
    body: list[str] = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            body.append("")
            continue
        if stripped.startswith("# "):
            title = stripped[2:].strip()
        elif stripped.startswith("## "):
            body.append(stripped[3:].strip())
        else:
            body.append(line)

    return title or "Untitled Slide", body


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color, radius=False):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    return box


def add_footer(slide, page_num: int, total: int):
    add_rect(slide, Inches(0.65), Inches(6.88), Inches(12.0), Inches(0.02), LINE)

    left_box = add_textbox(slide, Inches(0.7), Inches(6.92), Inches(4.4), Inches(0.25))
    p = left_box.text_frame.paragraphs[0]
    p.text = "Sarvam AI brand teardown | secondary research strategy deck"
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED

    right_box = add_textbox(slide, Inches(11.85), Inches(6.92), Inches(0.7), Inches(0.25))
    p = right_box.text_frame.paragraphs[0]
    p.text = f"{page_num}/{total}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED


def add_title(slide, title: str):
    title_box = add_textbox(slide, Inches(0.9), Inches(0.55), Inches(11.4), Inches(0.75))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT
    add_rect(slide, Inches(0.9), Inches(1.34), Inches(1.45), Inches(0.06), ACCENT)


def font_size_for_body(lines: list[str]) -> int:
    content_count = len([line for line in lines if line.strip()])
    if content_count > 14:
        return 13
    if content_count > 11:
        return 15
    if content_count > 8:
        return 17
    return 19


def add_body(slide, lines: list[str]):
    box = add_textbox(slide, Inches(1.0), Inches(1.65), Inches(11.3), Inches(4.95))
    frame = box.text_frame
    frame.clear()

    body_font = Pt(font_size_for_body(lines))
    first = True
    for raw_line in lines:
        if first:
            p = frame.paragraphs[0]
            first = False
        else:
            p = frame.add_paragraph()

        stripped = raw_line.strip()

        if not stripped:
            p.text = ""
            p.space_after = Pt(5)
            continue

        level = 0
        text = stripped
        if stripped.startswith("- "):
            level = 0
            text = stripped[2:].strip()
        elif raw_line.startswith("  - "):
            level = 1
            text = raw_line.strip()[2:].strip()
        elif raw_line.startswith("    - "):
            level = 2
            text = raw_line.strip()[2:].strip()

        p.text = text
        p.level = level
        p.font.size = body_font
        p.font.color.rgb = TEXT if level == 0 else MUTED
        p.space_after = Pt(7 if level == 0 else 4)
        p.line_spacing = 1.05

        if level == 0 and text.endswith(":"):
            p.font.bold = True
            p.font.color.rgb = ACCENT
        elif level == 0 and text.count(":") == 1 and len(text) < 42:
            p.font.bold = True


def add_cover_slide(prs: Presentation, title: str, body_lines: list[str], page_num: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BG)

    add_rect(slide, Inches(0.0), Inches(0.0), Inches(13.333), Inches(0.26), ACCENT)
    add_rect(slide, Inches(10.7), Inches(0.0), Inches(2.633), Inches(0.26), ACCENT_2)
    add_rect(slide, Inches(0.7), Inches(0.9), Inches(0.18), Inches(4.7), ACCENT)
    add_rect(slide, Inches(0.92), Inches(0.9), Inches(0.05), Inches(4.7), ACCENT_2)

    title_box = add_textbox(slide, Inches(1.2), Inches(1.0), Inches(8.8), Inches(1.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = TEXT

    y = Inches(2.35)
    if body_lines:
        subtitle = body_lines[0].strip()
        if subtitle:
            sub_box = add_textbox(slide, Inches(1.2), y, Inches(8.9), Inches(0.7))
            p = sub_box.text_frame.paragraphs[0]
            p.text = subtitle
            p.font.size = Pt(18)
            p.font.color.rgb = ACCENT
            y += Inches(0.55)

    details = add_textbox(slide, Inches(1.2), y, Inches(9.3), Inches(2.5))
    frame = details.text_frame
    frame.clear()
    first = True
    for line in body_lines[1:]:
        if not line.strip():
            continue
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        text = line.strip()
        if text.startswith("- "):
            text = text[2:].strip()
        p.text = text
        p.font.size = Pt(16)
        p.font.color.rgb = MUTED
        p.space_after = Pt(9)

    add_rect(slide, Inches(9.95), Inches(1.0), Inches(2.55), Inches(3.25), BG_SOFT, radius=True)
    callout = add_textbox(slide, Inches(10.2), Inches(1.25), Inches(2.05), Inches(2.75))
    frame = callout.text_frame
    p = frame.paragraphs[0]
    p.text = "Focus"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT

    for line in [
        "Brand equity",
        "Positioning",
        "GTM operating model",
        "Future stretch",
    ]:
        p = frame.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT
        p.space_before = Pt(8)

    add_footer(slide, page_num, total)


def add_standard_slide(prs: Presentation, title: str, body_lines: list[str], page_num: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BG)

    add_rect(slide, Inches(0.0), Inches(0.0), Inches(13.333), Inches(0.14), BG_SOFT)
    add_rect(slide, Inches(0.0), Inches(0.0), Inches(2.4), Inches(0.14), ACCENT)
    add_rect(slide, Inches(0.64), Inches(0.58), Inches(0.08), Inches(5.95), ACCENT)

    add_title(slide, title)
    add_body(slide, body_lines)
    add_footer(slide, page_num, total)


def build_deck():
    markdown_text = SOURCE_MD.read_text(encoding="utf-8")
    sections = split_slides(markdown_text)
    parsed = [parse_slide(section) for section in sections]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total = len(parsed)

    for idx, (title, body_lines) in enumerate(parsed, start=1):
        if idx == 1:
            add_cover_slide(prs, title, body_lines, idx, total)
        else:
            add_standard_slide(prs, title, body_lines, idx, total)

    prs.save(str(OUTPUT_PPTX))
    print(f"Saved {OUTPUT_PPTX}")


if __name__ == "__main__":
    build_deck()
